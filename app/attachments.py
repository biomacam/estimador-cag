"""Local text extraction for PDF and Word attachments (no provider Files API).

Attachments are untrusted external input, exactly like the transcript: the text
extracted here is only concatenated to the transcript with a clear separator,
and the whole result still travels inside the untrusted-data boundary of the
prompt and through the input guardrails.
"""

from __future__ import annotations

import io
import re
from pathlib import PurePosixPath

import structlog
from docx import Document
from pypdf import PdfReader

log = structlog.get_logger()

SUPPORTED_EXTENSIONS = (".pdf", ".docx")


class AttachmentError(Exception):
    """Base class; ``filename`` is the sanitised client-supplied name."""

    def __init__(self, filename: str, message: str) -> None:
        super().__init__(message)
        self.filename = filename
        self.message = message


class UnsupportedAttachmentError(AttachmentError):
    """The file extension is not a supported PDF/Word format."""


class AttachmentTooLargeError(AttachmentError):
    """The raw upload exceeds the configured byte limit."""


class AttachmentExtractionError(AttachmentError):
    """The file is corrupt, encrypted or has no extractable text."""


def safe_filename(raw: str) -> str:
    """Basename without control characters, so it cannot break the separator line."""
    name = PurePosixPath(raw.replace("\\", "/")).name
    return re.sub(r"[\x00-\x1f\x7f]", "", name) or "attachment"


def extract_text(filename: str, content: bytes, *, max_chars: int) -> str:
    """Return the text of a PDF or DOCX, truncated to ``max_chars`` with a marker."""
    name = safe_filename(filename)
    extension = PurePosixPath(name.lower()).suffix
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedAttachmentError(
            name, f"Unsupported attachment type {extension or '(none)'!r}; use PDF or DOCX."
        )

    try:
        text = _extract_pdf(content, name) if extension == ".pdf" else _extract_docx(content)
    except AttachmentError:
        raise
    except Exception as exc:  # noqa: BLE001 — parser libraries raise many unrelated types
        log.warning("attachment_extraction_failed", filename=name, error_type=type(exc).__name__)
        raise AttachmentExtractionError(name, "The file could not be read; it may be corrupt.") from exc

    text = text.strip()
    if not text:
        raise AttachmentExtractionError(
            name, "No extractable text found (a scanned document needs OCR, which is not supported)."
        )
    if len(text) > max_chars:
        log.info("attachment_truncated", filename=name, chars=len(text), max_chars=max_chars)
        text = f"{text[:max_chars]}\n[attachment truncated at {max_chars} characters]"
    return text


def _extract_pdf(content: bytes, name: str) -> str:
    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        raise AttachmentExtractionError(name, "Encrypted PDFs are not supported.")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(parts)


def enrich_transcript(transcript: str, attachments: list[tuple[str, str]]) -> str:
    """Append each ``(filename, text)`` after the transcript under a clear separator."""
    sections = [transcript]
    sections.extend(f"--- attachment: {safe_filename(name)} ---\n{text}" for name, text in attachments)
    return "\n\n".join(sections)
