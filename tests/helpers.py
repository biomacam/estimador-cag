"""Shared test helpers for asserting the untrusted-data delimiter contract
(see app.services.llm_service.frame_untrusted_input) without hardcoding the
tag in assertions about the wrapped text.
"""

import io
import re

from docx import Document

from app.schemas.estimation import EstimationResult


def make_pdf(text: str) -> bytes:
    """Smallest valid one-page PDF whose only content is ``text`` (no parentheses)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    return bytes(out)


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def valid_estimation_result() -> EstimationResult:
    return EstimationResult(
        summary="Build a booking system with payments.",
        confidence_pct=80,
        phases=[{
            "name": "Implementation", "duration_weeks": 4, "cost_eur": 10000,
            "summary": "Build booking and payment integrations.",
        }],
        total_duration_weeks=4,
        total_cost_eur=10000,
    )

_FRAMED_RE = re.compile(r"^<user-data-([a-z0-9]+)>\n(.*)\n</user-data-\1>$", re.DOTALL)


def unwrap_untrusted_input(framed: str) -> str:
    """Strip the ``<user-data-{tag}>...</user-data-{tag}>`` wrapper and return the
    original text. Raises AssertionError if ``framed`` doesn't match that shape.
    """
    match = _FRAMED_RE.fullmatch(framed)
    assert match is not None, f"Expected a <user-data-...> framed string, got: {framed!r}"
    return match.group(2)
